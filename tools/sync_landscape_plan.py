"""Synchronize the editable plans and regenerate landscape design diagrams."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt
import re

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / '策划案配图'
font_path = next(p for p in Path('C:/Windows/Fonts').glob('*') if p.name.lower() == 'msyh.ttc')
def font(size): return ImageFont.truetype(str(font_path), size)
def canvas(title):
    im = Image.new('RGB', (1600, 1000), '#f3f0e7')
    draw = ImageDraw.Draw(im)
    draw.text((60, 35), title, font=font(38), fill='#223c3b')
    return im, draw
def box(d, xy, title, lines, fill='#e4e6db'):
    d.rounded_rectangle(xy, radius=12, fill=fill, outline='#8a805e', width=2)
    x,y,_,_=xy
    d.text((x+20,y+14),title,font=font(27),fill='#233d39')
    for i,line in enumerate(lines): d.text((x+20,y+56+i*35),line,font=font(23),fill='#40504a')

im,d=canvas('檐下千秋  横屏界面线框图  1280 × 720')
box(d,(60,110,780,490),'01  加载页 / 欢迎页',['全幅水墨背景 · 欢迎原画自带作品标题'])
box(d,(85,235,365,460),'加载',['独立 Logo','真实读取阶段进度','连接失败 → 重试'])
box(d,(390,235,755,460),'欢迎 · 右侧操作',['启程 / 继续旅程','选择关卡','声音设置 / 退出'])
box(d,(820,110,1540,490),'02  专门的关卡选择页',['顶部：古建行旅 / 总进度 / 返回欢迎'])
box(d,(845,210,1515,295),'01 廊桥 → 02 社庙 → 03 戏楼 → 04 牌坊 → 05 墨塔',[], '#d6dece')
box(d,(845,315,1160,465),'场景预览',['使用当前章节原画'])
box(d,(1180,315,1515,465),'章节任务',['已修复 / 可进入','未解锁 + 原因'])
box(d,(60,535,780,935),'03  调查 / 修复',['顶部：识海 · 墨痕 · 侵蚀进度'])
box(d,(85,635,485,840),'左侧宽幅场景',['调查热点 / 描摹 / 构件','画布保持可见'])
box(d,(505,635,755,840),'右侧说明',['线索 / 任务目标','独立纵向滚动'])
d.text((90,872),'底栏：关卡 · 心舍 · 账册 · 设置 · 剧情',font=font(25),fill='#233d39')
box(d,(820,535,1540,935),'04  守护战',['左侧战场 / 右侧战斗说明'])
box(d,(845,635,1515,785),'固定技能栏',['1 挥墨  /  2 斗拱  /  3 藻井  /  4 飞檐  /  5 闪身'])
d.text((845,818),'准备按钮与技艺按钮不随说明滚动',font=font(25),fill='#233d39')
d.text((845,865),'仅显示已习得技艺；底部保留全局导航',font=font(25),fill='#233d39')
im.save(FIG/'图表-横屏UI线框图.png')

im,d=canvas('檐下千秋  页面结构与旅程流转')
box(d,(60,130,485,285),'加载',['连接本地服务 → 读取存档','读取章节 → 恢复会话'])
box(d,(590,130,1015,285),'欢迎',['启程 / 继续 / 关卡','声音设置 / 退出'])
box(d,(1120,130,1540,285),'关卡选择',['五章顺序导航','预览 · 锁定原因 · 继续'])
d.text((510,183),'→',font=font(40),fill='#775c36'); d.text((1040,183),'→',font=font(40),fill='#775c36')
box(d,(60,370,485,555),'开场对话 / 调查',['对话可返回关卡','三处调查 → 开始修复'])
box(d,(590,370,1015,555),'修复 / 记忆取舍',['描摹 · 构件 · 技艺 · 图式','保留墨灵 / 满额时替换'])
box(d,(1120,370,1540,555),'守护 / 结算',['八场守护战 / 两任务直接结算','结语 → 下一任务 / 关卡'])
d.text((510,427),'→',font=font(40),fill='#775c36'); d.text((1040,427),'→',font=font(40),fill='#775c36')
d.line((1330,285,1330,330,270,330,270,370),fill='#8a805e',width=3)
box(d,(60,650,775,895),'辅助页面',['心舍 → 墨灵详情 / 对照；账册；剧情回顾；设置','调查、解谜、战斗暂停后可继续','返回欢迎不清档；重启按服务端会话恢复'])
box(d,(825,650,1540,895),'失败支线',['失败演出：飞鸟山 → 回溯操作','恢复本次事件；保留已确认记忆与取舍','加载错误：提示原因 → 重新加载 / 退出'])
im.save(FIG/'图表-横屏页面结构图.png')

formal = ROOT/'Task01/Original/正式成果/《檐下千秋》游戏策划案（正式稿）.md'
dev = ROOT/'LLM-tmp/01-策划案/剧情任务功能策划案.md'
replacements = {
 '540×960 竖屏，默认窗口 450×800':'1280×720 横屏，默认窗口 1280×720',
 '540×960 逻辑竖屏，默认窗口 450×800':'1280×720 逻辑横屏，默认窗口 1280×720',
 '宽屏居中保留约 540 逻辑像素阅读区':'按横屏画布伸缩，宽屏扩展可用空间，最小窗口 960×540',
 '界面采用**古卷轴**形态，正文区为可纵向滚动的"谱页"':'界面采用**横屏水墨行旅**布局，调查、描摹、构件与战斗画布位于左侧，说明位于右侧独立滚动区',
 '界面采用古卷轴形态，正文区为可纵向滚动的"谱页"':'界面采用横屏水墨行旅布局，调查、描摹、构件与战斗画布位于左侧，说明位于右侧独立滚动区',
 '打开游戏后直接进入地图':'打开游戏后先经过加载页与欢迎页，再选择关卡或继续旅程',
 '纵向滚动"谱页"；标签 18 号，自动换行':'横屏场景与说明分栏；说明可独立滚动；标签 18 号，自动换行',
 '五列常驻：地图 · 心舍 · 账册 · 设置 · 剧情':'进入旅程后五列常驻：关卡 · 心舍 · 账册 · 设置 · 剧情',
 '两列技艺按钮':'五列横向技艺按钮',
 '两列技艺栏':'五列横向技艺栏',
 '中间滚动区承载任务正文与场景':'中间以左侧画布与右侧独立滚动说明承载调查、修复和战斗',
 '底部固定地图、心舍、账册、设置、剧情五个入口':'进入旅程后底部固定关卡、心舍、账册、设置、剧情五个入口',
 '游戏共 **16 个页面态**，主链为纵向顺序，底部五个入口在加载完成后常驻':'游戏共 **17 个页面态**，新增独立欢迎页，主链为加载、欢迎、关卡与事件；底部五个入口在进入旅程后常驻',
 '游戏共 16 个页面态，主链为纵向顺序，底部五个入口在加载完成后常驻':'游戏共 17 个页面态，新增独立欢迎页，主链为加载、欢迎、关卡与事件；底部五个入口在进入旅程后常驻',
 '开屏 / 古建地图':'加载 / 欢迎 / 关卡选择',
 '16 个页面态、页面流转':'17 个页面态、页面流转',
 '下图为四个核心界面的线框设计，展示顶栏、正文区、底部导航与战斗操作栏的层级关系：① 古建地图、② 建筑调查、③ 白蚀守护战、④ 事件结算。':'下图展示横屏加载与欢迎、关卡选择、调查与修复、守护战的分区关系。分辨率以 1280×720 为基准；欢迎与关卡使用独立全幅原画。',
 '图表-UI线框图.png':'图表-横屏UI线框图.png',
 '图表-页面结构图.png':'图表-横屏页面结构图.png',
}
spec = '''### 4.4 横屏入口与素材更新 2026年9月28日

本轮保留内容协议 v10、五章十任务和服务端权威存档，重新设计横屏入口与游戏内布局。启动顺序为“加载 → 欢迎 → 关卡选择 → 开场对话 → 调查 → 修复 → 取舍 → 守护／结算”。欢迎页不显示游戏内资源栏或导航。

加载页使用独立 Logo，按连接服务、读取旅程、展开章节、恢复会话显示真实阶段进度。只有全部步骤成功才进入欢迎页；失败显示原因、重新加载与退出，不建立虚假的离线进度。欢迎背景自带题字，操作集中于右下方，提供启程或继续旅程、选择关卡、声音设置和退出。

关卡页上方横排五章，依次为廊桥、社庙、戏楼、牌坊、墨塔；章节按钮显示修复数量或未解锁状态，选中章节用暖金边框强调。下方左侧展示章节场景，右侧仅展示该章任务与解锁原因，避免十任务长列表。可预览锁定章节，但任务按钮不可进入；已修复任务标注重访；存在未完事件时显示继续入口。返回时保留所选章节。

调查、描摹、构件修复、技艺机关和战斗采用左侧宽画布、右侧可滚动说明；图式选择保持横向并列。技能栏在底部横排，滚动说明不带走技能按钮。对话采用全幅背景与底部横向台词板。收藏、账册和结算长内容保持可滚动。窗口支持 960×540、1280×720、1600×900 与宽屏扩展。

素材来源为 Task02/Original-new/生成图片：游戏开始界面、关卡选择界面背景、独立 Logo，以及十任务场景与飞鸟山共十一张横屏场景。运行时副本集中于客户端 assets/landscape；欢迎背景中的主角使用新版原画，游戏内可移动角色仍为墨色剪影。Logo 原图为带浅色棋盘底的 RGB 图片，界面以着色器去除浅底显示墨字，保留源文件。场景侵蚀仍由褪色与裂纹叠层表现。

继续旅程优先恢复当前事件；首次启程进入关卡页。返回欢迎不清档、不新建玩家，不改变章节解锁、任务顺序、记忆取舍或战斗结算规则。未开始修复的调查在本次运行内暂停保留；关闭程序后，未提交的调查线索需要重新发现，已提交的会话按服务端进度恢复。

'''
def update(text):
    for a,b in replacements.items(): text=text.replace(a,b)
    return text
text=update(formal.read_text(encoding='utf-8'))
idx=text.index('\n## 五、')
if '### 4.4 横屏入口与素材更新' not in text:
    text=text[:idx]+'\n'+spec+text[idx:]
formal.write_text(text.rstrip()+'\n',encoding='utf-8')
text=update(dev.read_text(encoding='utf-8'))
if '## 横屏入口与素材更新' not in text:
    text+='\n'+spec.replace('### 4.4','##')
dev.write_text(text.rstrip()+'\n',encoding='utf-8')

for path in [formal.with_suffix('.docx'), ROOT/'LLM-tmp/01-策划案/《檐下千秋》游戏策划案.docx']:
    doc=Document(path)
    paragraphs=list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells: paragraphs.extend(cell.paragraphs)
    for p in paragraphs:
        value=update(p.text)
        if value!=p.text:
            # Keep paragraph style and first run's typography for changed prose.
            if p.runs:
                p.runs[0].text=value
                for run in p.runs[1:]: run.text=''
            else: p.add_run(value)
    if not any(p.text=='横屏入口与素材更新 2026年9月28日' for p in doc.paragraphs):
        heading=doc.add_paragraph('横屏入口与素材更新 2026年9月28日')
        for run in heading.runs:
            run.bold=True
            run.font.size=Pt(16)
        heading.paragraph_format.page_break_before=True
        for para in spec.split('\n\n')[1:]:
            if para.strip():
                p=doc.add_paragraph(para.strip())
                p.paragraph_format.line_spacing=1.5
                for run in p.runs:
                    run.font.size=Pt(12)
    # Replace the existing design diagrams by relationship order; original
    # document contains exactly the UI, page flow, gameplay diagrams.
    shapes=list(doc.inline_shapes)
    if path==formal.with_suffix('.docx'):
        for shape, filename in zip(shapes[:2], ['图表-横屏UI线框图.png','图表-横屏页面结构图.png']):
            rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            doc.part.related_parts[rid]._blob=(FIG/filename).read_bytes()
            shape.height=int(shape.width*1000/1600)
    doc.save(path)
print('Updated formal/development Markdown and Word plans; regenerated two diagrams.')
